"""สร้าง "พจนานุกรมข้อมูล" (data dictionary) ของฐานข้อมูล PostgreSQL เป็นไฟล์ Excel

ใช้ส่งให้คนภายนอกที่จะต่อเข้ามาอ่านฐานข้อมูลเอง (คู่กับ deploy/db_guide.md + external_analyst.md)

ข้อมูลมาจาก 2 ที่:
  1. **โครงสร้างจริงบนเซิร์ฟเวอร์** (ตาราง/คอลัมน์/ชนิด/คีย์/จำนวนแถว/ขนาด/ช่วงวันที่)
     อ่านผ่าน `ssh oxlet psql` ด้วยบัญชี `claude` (อ่านอย่างเดียว) — นับแถว + อ่านโครงสร้างเท่านั้น
     **ไม่ดึงเนื้อข้อมูลออกมาเลย**
  2. **คำอธิบายจากโค้ด** — คำอธิบายตารางจาก `dashboard/services/db_inventory.TABLES`
     (แหล่งเดียวกับหน้าสารบัญในเว็บ) + ชื่อไทยของคอลัมน์จาก `verbose_name` ในโมเดล

รัน:
    python scripts\\gen_db_dictionary.py [--out deploy\\db_dictionary.xlsx]

    ต้องใช้ Python ที่มีทั้ง Django และ openpyxl (openpyxl ไม่อยู่ใน requirements.txt เพราะ
    เซิร์ฟเวอร์ไม่ต้องใช้ — Python หลักของเครื่อง dev มีครบ) + ssh alias `oxlet` ใน ~/.ssh/config

เพิ่มตารางใหม่ → เติมคำอธิบายใน db_inventory.TABLES แล้วรันสคริปต์นี้ใหม่
(ตารางที่ยังไม่มีคำอธิบายจะขึ้นสีเหลืองว่า "ยังไม่มีคำอธิบาย")
"""
import argparse
import collections
import os
import subprocess
import sys
from datetime import datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
os.chdir(ROOT)
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "oxlet.settings")
os.environ["DB_HOST"] = ""          # ไม่ต้องต่อ Postgres ในเครื่อง (อ่านแค่โมเดล)
os.environ["DEBUG"] = "True"

import django  # noqa: E402

django.setup()

from django.apps import apps  # noqa: E402
try:
    from openpyxl import Workbook  # noqa: E402
except ImportError:
    raise SystemExit("ต้องมี openpyxl — รันด้วย Python หลักของเครื่อง หรือ pip install openpyxl")
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from dashboard.services import db_inventory as INV  # noqa: E402

SSH_HOST = "oxlet"
DB = "oxlet"


# ───────────────────────── อ่านจากเซิร์ฟเวอร์ ─────────────────────────

def psql(sql):
    """รัน SQL บนเซิร์ฟเวอร์ คืน list ของแถว (คั่นช่องด้วย \\0 กันข้อความมี tab/|)"""
    r = subprocess.run(["ssh", SSH_HOST, f"psql -d {DB} -qAt -z -v ON_ERROR_STOP=1"],
                       input=sql, capture_output=True, text=True, encoding="utf-8")
    if r.returncode != 0:
        raise SystemExit(f"psql ล้ม: {r.stderr.strip()}")
    return [ln.split("\0") for ln in r.stdout.splitlines() if ln]


def fetch_server():
    tables = psql("""
        select c.relname, c.relkind, pg_total_relation_size(c.oid),
               coalesce(array_to_string(c.relacl, ' '), '')
        from pg_class c join pg_namespace n on n.oid = c.relnamespace
        where n.nspname = 'public' and c.relkind in ('r', 'v') order by 1;""")
    cols = psql("""
        select table_name, ordinal_position, column_name, data_type,
               coalesce(character_maximum_length::text, ''), is_nullable
        from information_schema.columns where table_schema = 'public'
        order by table_name, ordinal_position;""")
    cons = psql("""
        select cl.relname, a.attname, co.contype, coalesce(fcl.relname, ''), coalesce(fa.attname, '')
        from pg_constraint co
        join pg_class cl on cl.oid = co.conrelid
        join pg_namespace n on n.oid = cl.relnamespace and n.nspname = 'public'
        join lateral unnest(co.conkey) with ordinality k(attnum, ord) on true
        join pg_attribute a on a.attrelid = cl.oid and a.attnum = k.attnum
        left join pg_class fcl on fcl.oid = co.confrelid
        left join pg_attribute fa on fa.attrelid = co.confrelid and fa.attnum = co.confkey[k.ord]
        where co.contype in ('p', 'f', 'u') order by 1, 2;""")
    kv = psql("""
        select case when key like '%:%' then split_part(key, ':', 1) || ':…' else key end,
               count(*), sum(length(data::text)), max(updated_at)::date
        from dash_kv group by 1 order by 3 desc;""")
    ver = psql("select current_setting('server_version');")[0][0].split()[0]
    return tables, cols, cons, kv, ver


# คอลัมน์เวลาที่ใช้บอก "ข้อมูลตั้งแต่-ถึง" (ไล่ตามลำดับ เจอตัวแรกใช้ตัวนั้น)
TIME_PREF = ["sent_at", "at", "snap_date", "date", "date_iso", "checked_out_at", "week_start",
             "first_seen", "fetched_at", "received_at", "connected_at", "created_at", "last_seen",
             "action_time", "applied", "updated_at", "date_joined"]


def fetch_counts(col_map, base_tables):
    parts = []
    for t in base_tables:
        names = [c["name"] for c in col_map[t]]
        tc = next((p for p in TIME_PREF if p in names), "")
        if tc:
            parts.append(f"select '{t}', count(*), '{tc}', coalesce(min({tc})::date::text, ''), "
                         f"coalesce(max({tc})::date::text, '') from {t}")
        else:
            parts.append(f"select '{t}', count(*), '', '', '' from {t}")
    out = {}
    # นับวันตามเวลาไทย (เซิร์ฟเวอร์ตั้งเป็น UTC — ไม่ตั้ง ข้อมูลช่วงตี 0-7 โมงไทยจะตกไปเป็นวันก่อนหน้า)
    sql = "SET TIME ZONE 'Asia/Bangkok';\n" + "\nunion all\n".join(parts) + ";"
    for t, n, tc, mn, mx in psql(sql):
        out[t] = dict(rows=int(n), tcol=tc, mn=mn, mx=mx)
    return out


# ───────────────────────── คำอธิบาย ─────────────────────────

GROUPS = [
    ("โซเชียล & โฆษณา", lambda t: t.startswith(("dash_meta", "dash_tiktok", "dash_youtube"))
     or t in ("dash_social_daily", "dash_ads_daily")),
    ("แดชบอร์ดขาย & ระบบ", lambda t: t.startswith("dash_")),
    ("สต็อกรถ & สเตปงาน", lambda t: t.startswith("cars_")),
    ("พนักงาน & เช็คชื่อ", lambda t: t in ("checkout_employee", "checkout_checkin")),
    ("เบิก-คืนรถ", lambda t: t in ("checkout_carmovement", "checkout_movementphoto",
                                  "checkout_checklistconfig", "checkout_checklistitem",
                                  "checkout_violationlog", "checkout_equipmentissue")),
    ("แชท & ลูกค้า (CRM)", lambda t: t.startswith("checkout_")),
    ("ตารางระบบ Django", lambda t: True),
]


def group_of(t):
    for i, (name, fn) in enumerate(GROUPS):
        if fn(t):
            return i, name


# ตารางที่ไม่มีโมเดลของเรา (Django สร้างเอง)
EXTRA_TABLES = {
    "auth_user_groups": dict(name="ผู้ใช้ ↔ บทบาท", what="ใครอยู่บทบาทไหน (ตารางเชื่อม)", pii=False, keep=""),
    "auth_group_permissions": dict(name="บทบาท ↔ สิทธิ์ย่อย", what="ระบบสร้างเอง", pii=False, keep=""),
    "auth_user_user_permissions": dict(name="ผู้ใช้ ↔ สิทธิ์ย่อย", what="ระบบสร้างเอง", pii=False, keep=""),
    "django_migrations": dict(name="ประวัติ migrate",
                              what="ไล่ว่าอัปโครงฐานข้อมูลถึงขั้นไหนแล้ว (ใช้ดูว่าตารางใหม่เพิ่มเข้ามาเมื่อไหร่)",
                              pii=False, keep=""),
}

# สถานะพิเศษของตาราง (ที่ตัวเลขอย่างเดียวบอกไม่ได้)
TABLE_NOTE = {
    "checkout_coachlog": "ยังว่างบนเซิร์ฟเวอร์ — รอนำเข้าประวัติห้องโค้ช (ระหว่างนี้อยู่ใน Google Sheets แท็บ 'บันทึกโค้ช')",
    "checkout_chatowner": "เปิดใช้ 3 ต.ค.69 (หน้า Connect) — แถวลูกค้าเก่าถูกสร้างพร้อมกันวันนั้น",
    "checkout_chatownerlog": "เปิดใช้ 3 ต.ค.69 — เริ่มมีประวัติเมื่อเซลล์รับ/ตอบลูกค้า",
    "checkout_lineeventlog": "ยังไม่เปิดใช้",
    "checkout_violationlog": "ยังไม่เปิดใช้",
    "checkout_equipmentissue": "ยังไม่เปิดใช้",
    "checkout_checklistconfig": "ยังไม่เปิดใช้ (เช็คลิสต์ตอนนี้ฝังในโค้ด)",
    "checkout_checklistitem": "ยังไม่เปิดใช้ (เช็คลิสต์ตอนนี้ฝังในโค้ด)",
    "dash_form": "ยังไม่มีการส่งฟอร์ม",
    "dash_tiktok_event": "ยังไม่มี event เข้ามา",
    "django_session": "ว่างเสมอ — ระบบเก็บ session ในคุกกี้ ไม่ได้เก็บลงตาราง",
    "cars_car": "ช่วงวันที่นับจากวันที่สร้างแถว (นำเข้าจาก Car Spend) — วันรับรถเข้าจริงดูคอลัมน์ date_in",
    "checkout_movementphoto": "รูปตัวจริงอยู่ใน Google Drive — ตารางนี้เก็บแค่ id ที่ชี้ไป",
    "dash_social_daily": "คำนวณมาจากตาราง *_snapshot — ลบแล้วสร้างใหม่ได้เสมอ",
    "dash_ads_daily": "คำนวณมาจาก dash_meta_ad_daily — ลบแล้วสร้างใหม่ได้เสมอ",
    "dash_kv": "ที่เก็บของจุกจิกของระบบ — ดูแท็บ 'คีย์ใน dash_kv'",
}

# คำอธิบายคอลัมน์ที่โมเดลไม่ได้ตั้งชื่อไทยไว้ (หรือตั้งไว้แต่อ่านไม่รู้เรื่องสำหรับคนนอก)
COL_OVR = {
    ("*", "id"): "เลขลำดับแถว (ระบบใส่ให้)",
    ("*", "created_at"): "เวลาที่สร้างแถวนี้",
    ("*", "updated_at"): "เวลาที่แก้แถวนี้ล่าสุด",
    ("*", "car_id"): "รถคันไหน (รหัสรถ → cars_car.code)",
    ("*", "config_id"): "ชุดเช็คลิสต์ (→ checkout_checklistconfig.id)",
    ("*", "movement_id"): "รอบเบิก-คืน (→ checkout_carmovement.id)",
    ("*", "checklist_item_id"): "ข้อในเช็คลิสต์ (→ checkout_checklistitem.id)",
    ("*", "trigger"): "ดึงเพราะอะไร — cron = รอบอัตโนมัติเที่ยงคืน (ใช้คิดยอดรายวัน) · manual = กดดึงเองกลางวัน (ห้ามเอามาคิดยอดรายวัน)",
    ("dash_kv", "key"): "ชื่อคีย์ (ความหมายแต่ละคีย์ดูแท็บ 'คีย์ใน dash_kv')",
    ("dash_kv", "data"): "ค่าที่เก็บ (JSON) — บางคีย์มี LINE user id ของพนักงาน",
    ("dash_form", "kind"): "ชนิดฟอร์ม — finance = เช็คไฟแนนซ์ก่อนเซ็น · loan = ขอสินเชื่อ",
    ("dash_form", "data"): "ข้อมูลในฟอร์มทั้งก้อน (JSON)",
    ("cars_car", "code"): "รหัสรถ (คีย์หลัก — ตารางอื่นอ้างถึงด้วยค่านี้)",
    ("cars_car", "branch"): "รหัสสาขา (ตรงกับ cars_branch.code — ไม่ได้ผูกเป็น foreign key)",
    ("cars_car", "extra"): "ข้อมูลเพิ่มเติมที่นำเข้าจาก Car Spend (JSON) — ราคานำเข้า · ลิงก์รูป · รูปขาย",
    ("cars_presence", "identity"): "บัญชี (ชื่อผู้ใช้ หรือ LINE user id)",
    ("cars_presence", "name"): "ชื่อที่โชว์",
    ("cars_presence", "role"): "บทบาท",
    ("cars_presence", "page"): "อยู่หน้าไหนของเว็บ",
    ("cars_presence", "last_seen"): "เห็นล่าสุดเมื่อ",
    ("checkout_groupchat", "sticker_id"): "รหัสสติกเกอร์",
    ("checkout_groupchat", "sticker_package"): "รหัสชุดสติกเกอร์",
    ("checkout_groupchat", "channel"): "บอทตัวไหนได้ยินข้อความนี้ (LINE user id ของคนเดียวกันต่างกันตามบอท)",
    ("checkout_fbchat", "message_id"): "รหัสข้อความของ Facebook (ไม่ซ้ำ)",
    ("checkout_lineeventlog", "group_id"): "LINE group id",
    ("auth_user", "username"): "ชื่อบัญชี — คนที่เข้าผ่าน LINE จะเป็น line_<LINE user id>",
    ("auth_user", "password"): "แฮชรหัสผ่าน (ไม่ใช่รหัสจริง) — ไม่มีประโยชน์ต่อการวิเคราะห์",
    ("auth_permission", "content_type_id"): "ชนิดข้อมูล (→ django_content_type.id)",
    ("django_admin_log", "content_type_id"): "ชนิดข้อมูลที่ถูกแก้ (→ django_content_type.id)",
    ("django_admin_log", "object_repr"): "ชื่อของรายการที่ถูกแก้",
    ("django_admin_log", "action_flag"): "ทำอะไร — 1 = เพิ่ม · 2 = แก้ · 3 = ลบ",
    ("django_content_type", "app_label"): "ชื่อแอป",
    ("django_content_type", "model"): "ชื่อโมเดล",
    ("django_migrations", "app"): "ชื่อแอป",
    ("django_migrations", "name"): "ชื่อขั้น migrate",
    ("django_migrations", "applied"): "ทำเมื่อ",
    ("auth_user_groups", "user_id"): "ผู้ใช้ (→ auth_user.id)",
    ("auth_user_groups", "group_id"): "บทบาท (→ auth_group.id)",
    ("auth_group_permissions", "group_id"): "บทบาท (→ auth_group.id)",
    ("auth_group_permissions", "permission_id"): "สิทธิ์ย่อย (→ auth_permission.id)",
    ("auth_user_user_permissions", "user_id"): "ผู้ใช้ (→ auth_user.id)",
    ("auth_user_user_permissions", "permission_id"): "สิทธิ์ย่อย (→ auth_permission.id)",
    ("dash_social_daily", "platform"): "แพลตฟอร์ม — meta = Facebook · tiktok · youtube",
    ("dash_social_daily", "object_id"): "รหัสโพสต์/คลิป (ตรงกับ post_id / video_id ในตาราง snapshot ของแพลตฟอร์มนั้น)",
    ("dash_social_daily", "owner_id"): "รหัสเพจ/ช่อง (page_id / open_id / channel_id)",
    ("dash_seller_weekly", "rj"): "RJ (ลีดที่ถูกรีเจ็ก ไม่นับเป็นลีด)",
    ("cars_loginevent", "ip"): "IP ที่เข้ามา",
    ("checkout_movementphoto", "line_message_id"): "รหัสข้อความ LINE ของรูปนี้ (ถ้ามาจากกลุ่ม)",
    ("checkout_lineeventlog", "line_message_id"): "รหัสข้อความ LINE",
    ("dash_youtube_video_snapshot", "is_short"): "เป็น Shorts ไหม (ยาวไม่เกิน 3 นาที)",
}

# ข้อมูลอ่อนไหวรายคอลัมน์ — ป้ายสั้นๆ ให้คนอ่านรู้ว่าต้องระวังช่องไหน
LINE = "LINE user id"
NAME = "ชื่อคน"
CHAT = "บทสนทนา"
RAW = "ข้อมูลดิบ (อาจมีข้อมูลส่วนบุคคล)"
SECRET = "รหัสลับ — ไม่มีประโยชน์ต่อการวิเคราะห์"
SENS = {
    "checkout_groupchat": {"sender_id": LINE, "sender_name": NAME, "text": CHAT, "extra": RAW,
                           "sent_by_name": NAME},
    "checkout_fbchat": {"sender_id": "Facebook PSID", "sender_name": NAME, "text": CHAT, "extra": RAW,
                        "sent_by_name": NAME},
    "checkout_fbprofile": {"user_id": "Facebook PSID", "display_name": NAME, "nickname": NAME,
                           "inbox_link": "ลิงก์เปิดแชท", "raw": RAW},
    "checkout_lineprofile": {"user_id": LINE, "display_name": NAME, "status_message": "สเตตัสส่วนตัว",
                             "nickname": NAME, "raw": RAW},
    "checkout_employee": {"nickname": NAME, "display_name": NAME, "note": "หมายเหตุลา/ป่วย"},
    "checkout_checkin": {"user_id": LINE, "display_name": NAME, "full_address": "ที่อยู่/สถานที่",
                         "province": "ที่อยู่/สถานที่", "raw": RAW, "note": "หมายเหตุลา/ป่วย"},
    "checkout_customerneed": {"customer_name": NAME, "contact": "เบอร์/ช่องทางติดต่อ", "evidence": CHAT,
                              "note": "หมายเหตุ"},
    "checkout_carmovement": {"borrower_name": NAME, "borrower_line_id": LINE, "approved_by": NAME},
    "checkout_coachlog": {"who": NAME, "text": CHAT},
    "checkout_chatowner": {"last_preview": CHAT, "picture_url": "รูปโปรไฟล์"},
    "checkout_chatownerlog": {"emp_name": NAME, "by_name": NAME},
    "checkout_lineeventlog": {"sender_line_id": LINE, "text": CHAT, "raw": RAW},
    "checkout_violationlog": {"person": NAME, "person_line_id": LINE},
    "checkout_equipmentissue": {"reporter": NAME, "approved_by": NAME},
    "cars_loginevent": {"identity": LINE, "name": NAME, "ip": "IP / อุปกรณ์", "user_agent": "IP / อุปกรณ์"},
    "cars_presence": {"identity": LINE, "name": NAME},
    "cars_scanlog": {"worker_name": NAME, "worker_id": LINE},
    "auth_user": {"username": LINE, "password": SECRET, "first_name": NAME, "last_name": NAME,
                  "email": "อีเมล"},
    "dash_event_log": {"target": LINE, "detail": RAW},
    "dash_kv": {"data": RAW},
    "dash_form": {"data": "ข้อมูลลูกค้าในฟอร์ม"},
    "dash_tiktok_account": {"access_token": SECRET, "refresh_token": SECRET},
    "dash_tiktok_event": {"user_openid": "TikTok user id", "content": RAW, "raw": RAW},
    "dash_followup_log": {"seller": "ชื่อพนักงาน (ผลงานรายคน)"},
    "dash_seller_weekly": {"seller": "ชื่อพนักงาน (ผลงานรายคน)"},
    "django_session": {"session_data": SECRET},
}

# ความสัมพันธ์ที่ "ไม่ได้ผูกเป็น foreign key" แต่ใช้ join จริง
LOGICAL_JOINS = [
    ("cars_car", "branch", "cars_branch", "code", "รหัสสาขา (BK = บ้านเก่า · ON = อ่อนนุช)"),
    ("checkout_groupchat", "group_id", "checkout_linegroup", "group_id",
     "แยกแชทเป็นรายกลุ่ม + รู้ประเภทกลุ่ม (หรือใช้ view v_group_chat ที่ join ไว้ให้แล้ว)"),
    ("checkout_groupchat", "sender_id", "checkout_lineprofile", "user_id",
     "ผู้ส่ง → โปรไฟล์ · ⚠ คนเดียวมี id ต่างกันในแต่ละบอท ดูคอลัมน์ channel ประกอบ"),
    ("checkout_checkin", "user_id", "checkout_lineprofile", "user_id", "LINE id ของคนที่ส่งรูปเช็คชื่อ"),
    ("cars_scanlog", "worker_id", "checkout_lineprofile", "user_id", "คนที่เปลี่ยนสเตปรถ"),
    ("checkout_carmovement", "borrower_line_id", "checkout_lineprofile", "user_id", "คนที่เบิกรถ"),
    ("checkout_fbchat", "thread_id", "checkout_fbprofile", "thread_id", "ห้องแชท Facebook (เทียบคู่กับ channel = เพจ)"),
    ("checkout_fbchat", "sender_id", "checkout_fbprofile", "user_id",
     "PSID ผู้ส่ง (เทียบคู่กับ channel — PSID ออกต่อเพจ)"),
    ("dash_social_daily", "object_id", "dash_meta_post_snapshot", "post_id", "เมื่อ platform = meta"),
    ("dash_social_daily", "object_id", "dash_tiktok_video_snapshot", "video_id", "เมื่อ platform = tiktok"),
    ("dash_social_daily", "object_id", "dash_youtube_video_snapshot", "video_id", "เมื่อ platform = youtube"),
    ("dash_tiktok_video_snapshot", "open_id", "dash_tiktok_account", "open_id", "คลิปนี้อยู่ช่องไหน"),
    ("dash_tiktok_account_snapshot", "open_id", "dash_tiktok_account", "open_id", "ยอดของช่องไหน"),
    ("dash_youtube_video_snapshot", "channel_id", "dash_youtube_channel_snapshot", "channel_id", "คลิปนี้อยู่ช่องไหน"),
    ("dash_meta_post_snapshot", "page_id", "dash_meta_page_daily", "page_id", "โพสต์นี้อยู่เพจไหน"),
    ("dash_meta_ad_daily", "account_id", "dash_ads_daily", "account_id", "ผลรวมรายวันของบัญชีโฆษณา"),
    ("dash_followup_log", "seller", "checkout_employee", "nickname", "เทียบด้วยชื่อเล่น"),
    ("dash_seller_weekly", "seller", "checkout_employee", "nickname", "เทียบด้วยชื่อเล่น"),
    ("checkout_customerneed", "lead_code", "(Google Sheets) ชีตลีด", "Code",
     "รหัสลีด — ตัวลีดอยู่ใน Google Sheets ไม่ได้อยู่ในฐานข้อมูลนี้"),
    ("checkout_coachlog", "case_code", "(Google Sheets) ชีตลีด", "Code",
     "รหัสเคส — 1 แถวอาจมีหลายรหัสคั่นด้วยช่องว่าง"),
    ("checkout_movementphoto", "file", "(Google Drive)", "file id", "ตัวไฟล์รูปอยู่ใน Google Drive"),
    ("cars_scanlog", "media", "(Google Drive)", "file id", "รูป/วิดีโอที่แนบตอนเปลี่ยนสเตป"),
]

VIEW_INFO = {
    "v_group_chat": "แชทกลุ่ม LINE พร้อมชื่อกลุ่ม + ประเภทกลุ่ม · **ไม่มี LINE user id** (ปลอดภัยกว่าตารางดิบ) — กรองด้วย \"ประเภทกลุ่ม\"",
    "v_line_group": "สารบัญกลุ่ม LINE — ชื่อ · ประเภทงาน · จำนวนข้อความ · ล่าสุดเมื่อไหร่ · ไม่มีชื่อคน",
    "v_employees": "พนักงาน 1 แถว/คน พร้อม LINE user id ทุกบัญชี (สร้างไว้ให้ n8n)",
    "v_employee_line": "พนักงาน 1 แถว/บัญชี LINE (คนเดียวมีได้หลายบัญชี — บอทเก่า/บอทใหม่)",
}

QUESTIONS = [
    ("โพสต์/คลิปไหนคนดูเยอะสุด · ยอดรายวันของแต่ละชิ้น", "dash_social_daily", "มียอดรายวันพร้อมใช้ ไม่ต้องลบ snapshot เอง"),
    ("ยอดทั้งเพจ Facebook รายวัน (ที่ Facebook รายงานเอง)", "dash_meta_page_daily", ""),
    ("ค่าโฆษณา · ต้นทุนต่อแชท/ต่อลีด", "dash_ads_daily (รวม) · dash_meta_ad_daily (รายโฆษณา)",
     "ต้นทุน = sum(spend) / sum(chats) ห้ามเฉลี่ยค่าเฉลี่ย"),
    ("ผู้ติดตามของแต่ละช่อง TikTok / YouTube", "dash_tiktok_account_snapshot · dash_youtube_channel_snapshot", ""),
    ("มีรถอะไรในสต็อก ราคาเท่าไหร่ อยู่สเตปไหน", "cars_car (+ cars_branch)", "stage = show คือรถพร้อมขาย"),
    ("รถคันนี้ผ่านมือใครมาบ้าง เมื่อไหร่", "cars_scanlog", ""),
    ("ใครเบิกรถออกไป คืนหรือยัง", "checkout_carmovement (+ checkout_movementphoto)", "returned_at ว่าง = ยังไม่คืน"),
    ("พนักงานมีใคร ตำแหน่งอะไร เข้างานกี่โมง", "checkout_employee", "หรือ view v_employees"),
    ("ใครมาสาย มากี่วันในเดือนนี้", "checkout_checkin", "1 แถว = คน 1 วัน"),
    ("ลูกค้ากำลังหารถอะไร งบเท่าไหร่", "checkout_customerneed", "waiting = true คือยังรอรถที่ตรงสเปก"),
    ("มีกลุ่ม LINE อะไรบ้าง แต่ละกลุ่มทำงานอะไร", "checkout_linegroup", "หรือ view v_line_group"),
    ("แชทของกลุ่มใดกลุ่มหนึ่ง (เช่นห้องโค้ชเซลล์)", "view v_group_chat", "กรองด้วยคอลัมน์ \"ประเภทกลุ่ม\""),
    ("ลูกค้าคุยอะไรกับเรา", "checkout_groupchat (LINE) · checkout_fbchat (Facebook)", "แชทลูกค้าเก็บ 60 วัน"),
    ("เซลล์ตอบลูกค้าเร็วแค่ไหน · ลูกค้าเป็นของเซลล์คนไหน", "checkout_chatownerlog · checkout_chatowner",
     "เริ่มเก็บ 3 ต.ค.69 (หน้า Connect)"),
    ("ผลงานเซลล์รายสัปดาห์ (ลีด/จอง/ปล่อย)", "dash_seller_weekly", "ตัวเลขต้นทางอยู่ใน Google Sheets"),
    ("เซลล์โดนทวงงานเรื่องอะไรบ่อย", "dash_followup_log", ""),
    ("งานส่งข้อความอัตโนมัติล้มไหม เมื่อไหร่", "dash_event_log", "kind = 'line_send' and not ok"),
    ("ใครเข้าระบบเมื่อไหร่", "cars_loginevent", ""),
]

TRAPS = [
    ("ตาราง *_snapshot เก็บ \"ยอดสะสม\" ไม่ใช่ยอดรายวัน",
     "ยอดของวันนั้น = ค่าวันนี้ − ค่าเมื่อวาน · นับเฉพาะแถว trigger = 'cron' · วันเดียวกันมีหลายแถวได้ ให้หยิบแถวที่ดึงล่าสุด · ผลต่างติดลบไม่นับ · ต้องมีอย่างน้อย 2 วัน"),
    ("อยากได้ยอดรายวันเลย → ใช้ dash_social_daily",
     "คำนวณไว้แล้ว 1 แถว = โพสต์/คลิป 1 ชิ้น × 1 วัน · มีทั้งยอดที่เพิ่มวันนั้น (views…) และยอดสะสม (cum_…)"),
    ("ต้นทุนต่อหน่วย ต้องคิดจากผลรวมหารผลรวม",
     "sum(spend) / nullif(sum(chats), 0) — ห้าม avg(spend/chats) (วัดจริงต่างกัน 61.31 กับ 62.49 บาท)"),
    ("YouTube ไม่มียอดแชร์", "API ไม่ให้ → shares = 0 เสมอ ไม่ใช่บั๊ก"),
    ("ข้อมูลบางตารางมีวันหมดอายุ (PDPA)",
     "แชทกลุ่ม 90 วัน · แชทลูกค้า 60 วัน · ข้อมูลดิบจาก API 90 วัน · ล็อกเหตุการณ์ 90 วัน — นับย้อนหลังเกินนั้นไม่ได้ (ดูคอลัมน์ 'อายุข้อมูล' แท็บตาราง)"),
    ("LINE user id ของคนเดียวกันต่างกันตามบอท",
     "count(distinct sender_id) ≠ จำนวนคน — ดูคอลัมน์ channel ประกอบ หรือนับด้วย employee_id / ชื่อเล่น"),
    ("ตัวเลขการขาย (ลีด/จอง/ปล่อย) ไม่ได้อยู่ในฐานข้อมูลนี้",
     "อยู่ใน Google Sheets · ในนี้มีแค่สรุปรายสัปดาห์ (dash_seller_weekly) และสถิติการทวงงาน"),
]

NOT_IN_DB = [
    ("ข้อมูลการขาย · ลีด · ยอดจอง/ปล่อยรถ · เป้าเซลล์", "Google Sheets"),
    ("รูปรถ · วิดีโอ · เอกสาร", "Google Drive (ในฐานข้อมูลเก็บแค่ id ที่ชี้ไป)"),
    ("ไฟล์รูปในแชท LINE", "อยู่ที่ LINE (เก็บแค่ message_id + ธง has_media)"),
]

KV_EXTRA = {
    "line_oauth:…": "สถานะชั่วคราวระหว่าง login ด้วย LINE (หมดอายุ 10 นาที)",
    "tiktok_oauth:…": "ลิงก์ขออนุญาตเชื่อมช่อง TikTok (ใช้ได้ครั้งเดียว)",
    "checkin_group_member": "แคชว่าพนักงานคนไหนอยู่ในกลุ่มเช็คชื่อ",
    "checkin_notify_config": "ตั้งค่าส่งตารางเช็คชื่อเข้า LINE",
    "checkin_notify_last": "ส่งตารางเช็คชื่อล่าสุดเมื่อไหร่ (กันส่งซ้ำ)",
    "meta_page_names": "ชื่อเพจ Facebook",
    "meta_sync_last": "ผลการดึงข้อมูล Facebook รอบล่าสุด",
    "meta_sync_daily": "วันนี้ดึง Facebook แล้วหรือยัง",
    "meta_sync_lock": "ล็อกกันดึง Facebook ซ้อน",
    "tiktok_sync_last": "ผลการดึงข้อมูล TikTok รอบล่าสุด",
    "tiktok_sync_daily": "วันนี้ดึง TikTok แล้วหรือยัง",
    "tiktok_sync_lock": "ล็อกกันดึง TikTok ซ้อน",
    "youtube_sync_last": "ผลการดึงข้อมูล YouTube รอบล่าสุด",
    "youtube_sync_daily": "วันนี้ดึง YouTube แล้วหรือยัง",
    "youtube_sync_lock": "ล็อกกันดึง YouTube ซ้อน",
    "leadreport_plan": "แผนรายงานลีด",
    "cfg_sellers_config": "ตั้งค่าเซลล์ (เป้า/ทีม)",
    "cfg_schedule_config": "ตารางเวลาส่งแจ้งเตือน",
    "coach_sync_last": "ผลการซิงก์ห้องโค้ชรอบล่าสุด",
    "line_ingest_last": "ร่องรอยดีบักของข้อมูลที่ n8n ส่งมา",
    "purchase_fetch_last": "ผลการอ่านชีตจัดซื้อรอบล่าสุด",
    "chat_cleanup_last": "ลบแชทหมดอายุล่าสุดเมื่อไหร่",
    "eventlog_trim_last": "ตัดล็อกเหตุการณ์เก่าล่าสุดเมื่อไหร่",
    "cardline_lock": "ล็อกกันส่งการ์ดซ้อน",
}


def kv_label(k):
    if k in INV.KV_LABEL:
        return INV.KV_LABEL[k]
    if k in KV_EXTRA:
        return KV_EXTRA[k]
    if k.startswith("cardline_last_"):
        return f"ผลการส่งการ์ด {k[len('cardline_last_'):]} เข้า LINE รอบล่าสุด"
    if k.startswith("cardline_"):
        return f"ตั้งค่าการส่งการ์ด {k[len('cardline_'):]} เข้า LINE (เวลา/กลุ่ม/ช่วงวันที่)"
    return ""


def short_type(dt, ln):
    m = {"character varying": "varchar", "timestamp with time zone": "timestamptz",
         "timestamp without time zone": "timestamp", "double precision": "float8", "ARRAY": "array"}
    t = m.get(dt, dt)
    return f"{t}({ln})" if ln else t


def model_meta():
    out = {}
    for m in apps.get_models():
        fl = {}
        for f in m._meta.get_fields():
            if not getattr(f, "concrete", False) or f.many_to_many:
                continue
            help_ = str(getattr(f, "help_text", "") or "")
            if help_.isascii():
                help_ = ""
            fl[f.column] = dict(name=f.name, verbose=str(getattr(f, "verbose_name", "") or ""),
                                help=help_, choices=[(str(k), str(v)) for k, v in (f.choices or []) if str(k)])
        out[m._meta.db_table] = fl
    return out


def describe_col(table, col, meta):
    if (table, col) in COL_OVR:
        return COL_OVR[(table, col)]
    if col == "trigger":    # กับดักอันดับ 1 ของการคิดยอดรายวัน — อธิบายเต็มเสมอ ไม่ใช้ชื่อสั้นจากโมเดล
        return COL_OVR[("*", "trigger")]
    f = meta.get(table, {}).get(col)
    if f:
        v = f["verbose"]
        generic = (not v) or v.lower() in (f["name"].replace("_", " ").lower(), "id")
        if not generic:
            return v + (f" — {f['help']}" if f["help"] else "")
    if ("*", col) in COL_OVR:
        return COL_OVR[("*", col)]
    return ""


# ───────────────────────── เขียนไฟล์ ─────────────────────────

FONT = "Tahoma"
PURPLE = "5B21B6"
F_HEAD = Font(name=FONT, bold=True, color="FFFFFF", size=10)
F_BODY = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_MONO = Font(name="Consolas", size=10)
F_MONO_B = Font(name="Consolas", size=10, bold=True)
FILL_HEAD = PatternFill("solid", fgColor=PURPLE)
FILL_GROUP = PatternFill("solid", fgColor="EDE9FE")
FILL_PII = PatternFill("solid", fgColor="FFEDD5")
FILL_SAFE = PatternFill("solid", fgColor="DCFCE7")
FILL_WARN = PatternFill("solid", fgColor="FEF9C3")
FILL_MUTED = PatternFill("solid", fgColor="F3F4F6")
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)


def header(ws, row, titles, widths):
    for i, (t, w) in enumerate(zip(titles, widths), 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, CENTER, BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30


def put(ws, row, col, val, font=F_BODY, fill=None, align=WRAP, fmt=None):
    c = ws.cell(row=row, column=col, value=val)
    c.font, c.alignment, c.border = font, align, BORDER
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    return c


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, size=14, bold=True, color=PURPLE)
    ws["A2"] = sub
    ws["A2"].font = Font(name=FONT, size=10, italic=True, color="6B7280")


def build(out_path):
    print("อ่านโครงสร้างจากเซิร์ฟเวอร์…")
    tables, cols, cons, kv, ver = fetch_server()
    rel = {t: dict(kind=k, size=int(s), acl=a) for t, k, s, a in tables}
    col_map = collections.defaultdict(list)
    for t, pos, name, dt, ln, nul in cols:
        col_map[t].append(dict(pos=int(pos), name=name, type=short_type(dt, ln), null=(nul == "YES")))
    base = [t for t in rel if rel[t]["kind"] == "r"]
    views = [t for t in rel if rel[t]["kind"] == "v"]
    print(f"  ตาราง {len(base)} · view {len(views)} · นับแถว…")
    stats = fetch_counts(col_map, base)

    keys = collections.defaultdict(list)       # (t, col) → ["PK", "FK → x.y", "UNIQUE"]
    fks = []
    for t, c, typ, ft, fc in cons:
        if typ == "p":
            keys[(t, c)].append("PK")
        elif typ == "u":
            keys[(t, c)].append("UNIQUE")
        elif typ == "f":
            keys[(t, c)].append(f"FK → {ft}.{fc}")
            fks.append((t, c, ft, fc))

    meta = model_meta()
    info = dict(INV.TABLES)
    info.update(EXTRA_TABLES)
    ro_safe = {t for t in base if "ro_safe=" in rel[t]["acl"]}
    now = datetime.now()
    stamp = f"{now.day} {['', 'ม.ค.', 'ก.พ.', 'มี.ค.', 'เม.ย.', 'พ.ค.', 'มิ.ย.', 'ก.ค.', 'ส.ค.', 'ก.ย.', 'ต.ค.', 'พ.ย.', 'ธ.ค.'][now.month]} {now.year + 543}"

    order = sorted(base, key=lambda t: (group_of(t)[0], t))
    wb = Workbook()

    # ── แท็บ 2: ตาราง ──
    ws = wb.create_sheet("ตาราง")
    title(ws, "สารบัญตาราง — แต่ละตารางเก็บอะไร",
          f"สถานะ ณ {stamp} · จำนวนแถว/ช่วงวันที่นับจากฐานข้อมูลจริง · สีส้ม = มีข้อมูลส่วนบุคคล · สีเขียว = บัญชี ro_safe เห็นได้")
    H = ["#", "กลุ่ม", "ชื่อตาราง", "ชื่อภาษาไทย", "เก็บอะไร", "จำนวนแถว", "ขนาด (MB)",
         "ข้อมูลตั้งแต่", "ถึงวันที่", "นับช่วงวันที่จาก", "ข้อมูลส่วนบุคคล", "บัญชี ro_safe เห็น",
         "อายุข้อมูล / การลบ", "จำนวนคอลัมน์", "หมายเหตุ"]
    header(ws, 4, H, [5, 18, 30, 30, 60, 11, 10, 12, 12, 14, 12, 12, 30, 10, 45])
    r = 5
    first = r
    for i, t in enumerate(order, 1):
        d = info.get(t, {})
        st = stats[t]
        g = group_of(t)[1]
        missing = not d
        pii = bool(d.get("pii"))
        put(ws, r, 1, i, align=CENTER)
        put(ws, r, 2, g)
        put(ws, r, 3, t, font=F_MONO_B)
        put(ws, r, 4, d.get("name", "") or "ยังไม่มีคำอธิบาย", fill=FILL_WARN if missing else None)
        put(ws, r, 5, (d.get("what", "") or "").replace("**", ""))
        put(ws, r, 6, st["rows"], fmt="#,##0", align=Alignment(vertical="top"))
        put(ws, r, 7, round(rel[t]["size"] / 1048576, 2), fmt="#,##0.00", align=Alignment(vertical="top"))
        put(ws, r, 8, st["mn"], align=CENTER)
        put(ws, r, 9, st["mx"], align=CENTER)
        put(ws, r, 10, st["tcol"], font=F_MONO)
        put(ws, r, 11, "มี" if pii else "–", fill=FILL_PII if pii else None, align=CENTER)
        put(ws, r, 12, "เห็น" if t in ro_safe else "–", fill=FILL_SAFE if t in ro_safe else None, align=CENTER)
        put(ws, r, 13, d.get("keep", "") or "เก็บถาวร")
        put(ws, r, 14, len(col_map[t]), align=CENTER)
        note = TABLE_NOTE.get(t, "")
        if not note and st["rows"] == 0:
            note = "ยังไม่มีข้อมูล"
        put(ws, r, 15, note, fill=FILL_MUTED if st["rows"] == 0 else None)
        r += 1
    last = r - 1
    put(ws, r, 5, "รวมทั้งหมด", font=F_BOLD, align=Alignment(horizontal="right", vertical="top"))
    put(ws, r, 6, f"=SUM(F{first}:F{last})", font=F_BOLD, fmt="#,##0", align=Alignment(vertical="top"))
    put(ws, r, 7, f"=SUM(G{first}:G{last})", font=F_BOLD, fmt="#,##0.00", align=Alignment(vertical="top"))
    put(ws, r, 14, f"=SUM(N{first}:N{last})", font=F_BOLD, align=CENTER)
    total_row = r
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:O{last}"

    # ── แท็บ 3: คอลัมน์ ──
    wc = wb.create_sheet("คอลัมน์")
    title(wc, "รายละเอียดคอลัมน์ทุกตาราง",
          "ใช้ตัวกรองที่หัวตาราง (ช่อง 'ชื่อตาราง') เพื่อดูทีละตาราง · สีส้ม = ข้อมูลอ่อนไหว ต้องระวังเวลาใช้/ส่งต่อ")
    H = ["ชื่อตาราง", "ชื่อตาราง (ไทย)", "ลำดับ", "คอลัมน์", "ชนิดข้อมูล", "ความหมาย",
         "ค่าที่เป็นไปได้", "คีย์", "ว่างได้", "ข้อมูลอ่อนไหว"]
    header(wc, 4, H, [30, 26, 7, 24, 16, 48, 48, 30, 8, 22])
    r = 5
    for t in order:
        tname = info.get(t, {}).get("name", "")
        server_cols = {c["name"] for c in col_map[t]}
        rows = list(col_map[t])
        for extra_col in meta.get(t, {}):
            if extra_col not in server_cols:     # มีในโค้ดแต่ยังไม่ขึ้นเซิร์ฟเวอร์
                rows.append(dict(pos=None, name=extra_col, type="(ยังไม่มีบนเซิร์ฟเวอร์)", null=True))
        for c in rows:
            col = c["name"]
            f = meta.get(t, {}).get(col, {})
            choices = " · ".join(f"{k} = {v}" for k, v in f.get("choices", []))
            sens = SENS.get(t, {}).get(col, "")
            desc = describe_col(t, col, meta)
            if c["pos"] is None:
                desc = (desc + " — ").lstrip(" — ") + "เพิ่มในโค้ดแล้ว รอ deploy ขึ้นเซิร์ฟเวอร์"
            put(wc, r, 1, t, font=F_MONO)
            put(wc, r, 2, tname)
            put(wc, r, 3, c["pos"], align=CENTER)
            put(wc, r, 4, col, font=F_MONO_B)
            put(wc, r, 5, c["type"], font=F_MONO)
            put(wc, r, 6, desc, fill=FILL_WARN if not desc else None)
            put(wc, r, 7, choices)
            put(wc, r, 8, " · ".join(keys.get((t, col), [])), font=F_MONO)
            put(wc, r, 9, "ได้" if c["null"] else "–", align=CENTER)
            put(wc, r, 10, sens, fill=FILL_PII if sens else None)
            r += 1
    wc.freeze_panes = "E5"
    wc.auto_filter.ref = f"A4:J{r - 1}"
    n_cols = r - 5

    # ── แท็บ 4: ความสัมพันธ์ ──
    wr = wb.create_sheet("ความสัมพันธ์")
    title(wr, "ตารางไหนเชื่อมกับตารางไหน (ใช้ JOIN)",
          "FK = ผูกไว้ในฐานข้อมูลจริง · เทียบค่า = ไม่ได้ผูกไว้ แต่ค่าตรงกันและใช้ join ได้")
    header(wr, 4, ["จากตาราง", "คอลัมน์", "ไปที่ตาราง", "คอลัมน์", "ชนิด", "หมายเหตุ"], [30, 20, 32, 16, 12, 60])
    r = 5
    for t, c, ft, fc in sorted(fks, key=lambda x: (group_of(x[0])[0], x[0], x[1])):
        for i, v in enumerate([t, c, ft, fc], 1):
            put(wr, r, i, v, font=F_MONO)
        put(wr, r, 5, "FK", align=CENTER)
        put(wr, r, 6, "")
        r += 1
    for t, c, ft, fc, note in LOGICAL_JOINS:
        for i, v in enumerate([t, c, ft, fc], 1):
            put(wr, r, i, v, font=F_MONO)
        put(wr, r, 5, "เทียบค่า", align=CENTER, fill=FILL_WARN)
        put(wr, r, 6, note)
        r += 1
    wr.freeze_panes = "A5"
    wr.auto_filter.ref = f"A4:F{r - 1}"

    # ── แท็บ 5: View ──
    wv = wb.create_sheet("View")
    title(wv, "View สำเร็จรูป (ตารางเสมือนที่ join ไว้ให้แล้ว)",
          "ใช้ได้เหมือนตาราง: SELECT * FROM v_group_chat … · ใครเข้าได้ดูคอลัมน์ 'บัญชีที่อ่านได้'")
    header(wv, 4, ["View", "ใช้ทำอะไร", "บัญชีที่อ่านได้", "ลำดับ", "คอลัมน์", "ชนิดข้อมูล"], [20, 50, 22, 7, 26, 18])
    r = 5
    for v in sorted(views):
        roles = [x.split("=")[0] for x in rel[v]["acl"].split() if "=" in x and x.split("=")[0]]
        roles = [x for x in roles if x in ("ro_all", "ro_safe")] or ["(ภายในเท่านั้น)"]
        for i, c in enumerate(col_map[v]):
            put(wv, r, 1, v if i == 0 else "", font=F_MONO_B)
            put(wv, r, 2, VIEW_INFO.get(v, "").replace("**", "") if i == 0 else "")
            put(wv, r, 3, ", ".join(roles) if i == 0 else "", font=F_MONO)
            put(wv, r, 4, c["pos"], align=CENTER)
            put(wv, r, 5, c["name"], font=F_MONO)
            put(wv, r, 6, c["type"], font=F_MONO)
            r += 1

    # ── แท็บ 6: อยากรู้อะไร ──
    wq = wb.create_sheet("อยากรู้อะไร→ดูที่ไหน")
    title(wq, "อยากรู้อะไร → ดูตารางไหน", "จุดเริ่มต้นสำหรับคำถามที่เจอบ่อย")
    header(wq, 4, ["คำถาม", "ตาราง / view", "ข้อควรรู้"], [52, 50, 50])
    r = 5
    for q, t, n in QUESTIONS:
        put(wq, r, 1, q)
        put(wq, r, 2, t, font=F_MONO)
        put(wq, r, 3, n)
        r += 1

    # ── แท็บ 7: คีย์ใน dash_kv ──
    wk = wb.create_sheet("คีย์ใน dash_kv")
    title(wk, "ข้างในตาราง dash_kv มีอะไร",
          "dash_kv เป็นที่เก็บค่าตั้ง/สถานะของระบบ 1 แถวต่อคีย์ (ไม่สะสม) · ไม่ค่อยมีประโยชน์ต่อการวิเคราะห์ ยกเว้นคีย์ main")
    header(wk, 4, ["คีย์", "จำนวนแถว", "ขนาด (KB)", "อัปเดตล่าสุด", "ความหมาย"], [34, 10, 11, 13, 60])
    r = 5
    for k, n, size, upd in kv:
        put(wk, r, 1, k, font=F_MONO)
        put(wk, r, 2, int(n), fmt="#,##0", align=CENTER)
        put(wk, r, 3, round(int(size) / 1024, 1), fmt="#,##0.0", align=Alignment(vertical="top"))
        put(wk, r, 4, upd, align=CENTER)
        put(wk, r, 5, kv_label(k))
        r += 1
    wk.freeze_panes = "A5"

    # ── แท็บ 1: อ่านก่อน ──
    wa = wb.active
    wa.title = "อ่านก่อน"
    wa.column_dimensions["A"].width = 34
    wa.column_dimensions["B"].width = 90
    title(wa, "พจนานุกรมข้อมูล (Data Dictionary) — ฐานข้อมูล Oxlet Auto",
          f"สถานะ ณ {stamp} · PostgreSQL {ver} · ฐานข้อมูลชื่อ oxlet")
    r = 4

    def sec(text):
        nonlocal r
        r += 1
        c = wa.cell(row=r, column=1, value=text)
        c.font = Font(name=FONT, size=12, bold=True, color=PURPLE)
        r += 1

    def line(a, b, fill=None, fa=F_BOLD, fb=F_BODY):
        nonlocal r
        put(wa, r, 1, a, font=fa, fill=fill)
        put(wa, r, 2, b, font=fb, fill=fill)
        r += 1

    sec("ภาพรวม")
    line("จำนวนตาราง", f"=COUNTA('ตาราง'!C5:C{last})&\" ตาราง + {len(views)} view\"")
    line("จำนวนแถวรวม", f"=TEXT('ตาราง'!F{total_row},\"#,##0\")&\" แถว\"")
    line("ขนาดรวม", f"=TEXT('ตาราง'!G{total_row},\"#,##0.0\")&\" MB\"")
    line("ตารางที่มีข้อมูลส่วนบุคคล", f"=COUNTIF('ตาราง'!K5:K{last},\"มี\")&\" ตาราง (สีส้ม)\"")
    line("ตารางที่บัญชี ro_safe เห็น", f"=COUNTIF('ตาราง'!L5:L{last},\"เห็น\")&\" ตาราง (สีเขียว)\"")

    sec("แท็บในไฟล์นี้")
    for a, b in [
        ("ตาราง", "สารบัญทุกตาราง — เก็บอะไร · กี่แถว · ข้อมูลตั้งแต่เมื่อไหร่ · มีข้อมูลส่วนบุคคลไหม · เก็บนานแค่ไหน"),
        ("คอลัมน์", f"ทุกคอลัมน์ของทุกตาราง ({n_cols:,} คอลัมน์) — ชนิดข้อมูล · ความหมายภาษาไทย · ค่าที่เป็นไปได้ · คีย์ · ช่องที่อ่อนไหว"),
        ("ความสัมพันธ์", "ตารางไหน join กับตารางไหนด้วยคอลัมน์อะไร (รวมที่ไม่ได้ผูก FK แต่ใช้ join ได้)"),
        ("View", "ตารางเสมือนที่ join ไว้ให้แล้ว เช่น v_group_chat (แชทพร้อมชื่อกลุ่ม ไม่มี LINE id)"),
        ("อยากรู้อะไร→ดูที่ไหน", "คำถามที่เจอบ่อย → ตารางที่ต้องเปิด"),
        ("คีย์ใน dash_kv", "ข้างในตารางค่าตั้งของระบบ"),
    ]:
        line(a, b)

    sec("ป้ายสี")
    line("สีส้ม", "มีข้อมูลส่วนบุคคล (ชื่อคน · เบอร์ · LINE user id · บทสนทนา) — ใช้เท่าที่จำเป็น ห้ามก๊อปออกไปเก็บที่อื่น", fill=FILL_PII)
    line("สีเขียว", "บัญชีระดับ ro_safe (เห็นเฉพาะตัวเลขโซเชียล/โฆษณา/สต็อกรถ) อ่านตารางนี้ได้ · บัญชี ro_all อ่านได้ทุกตาราง", fill=FILL_SAFE)
    line("สีเหลือง", "เชื่อมด้วยการเทียบค่า (ไม่มี FK) หรือยังไม่มีคำอธิบาย", fill=FILL_WARN)
    line("สีเทา", "ตารางที่ยังไม่มีข้อมูล", fill=FILL_MUTED)

    sec("⚠ สิ่งที่ไม่ได้อยู่ในฐานข้อมูลนี้ (อ่านก่อนสรุปว่า \"ระบบเก็บแค่นี้\")")
    for a, b in NOT_IN_DB:
        line(a, b)

    sec("★ อ่านตัวเลขให้ถูก — กับดักที่เจอมาแล้วจริง")
    for a, b in TRAPS:
        line(a, b)

    sec("วิธีต่อเข้าฐานข้อมูล")
    for a, b in [
        ("ช่องทาง", "ฐานข้อมูลไม่เปิดออกอินเทอร์เน็ต — ต่อผ่านอุโมงค์ SSH เท่านั้น (คีย์ SSH + บัญชีฐานข้อมูลได้จากผู้ดูแลระบบ)"),
        ("ขั้นที่ 1 เปิดอุโมงค์", "ssh -i <ไฟล์คีย์> -N -L 15432:127.0.0.1:5432 <บัญชี ssh>@srv1793506.hstgr.cloud"),
        ("ขั้นที่ 2 ต่อฐานข้อมูล", "Host 127.0.0.1 · Port 15432 · Database oxlet · User/Password ตามที่ได้รับ · ไม่ต้องใช้ SSL"),
        ("เครื่องมือที่ใช้ได้", "psql · DBeaver · pgAdmin · Power BI · Excel · Metabase"),
        ("อ่านอย่างเดียว", "แก้/ลบข้อมูลไม่ได้ (ถูกปฏิเสธที่ระดับสิทธิ์)"),
        ("คำสั่งเกิน 30 วินาทีถูกตัด", "ใส่ WHERE จำกัดช่วงวันที่ หรือ LIMIT"),
        ("1 คน = 1 บัญชี", "ห้ามใช้ร่วมกัน/ส่งต่อ · ทุกคำสั่งมีชื่อบัญชีกำกับในล็อก"),
        ("เวลา", "เซิร์ฟเวอร์ตั้งโซนเวลาเป็น UTC — อยากเห็นเวลาไทยให้สั่ง  SET TIME ZONE 'Asia/Bangkok';  ก่อน query "
                 "(หรือใช้ col AT TIME ZONE 'Asia/Bangkok') · ช่วงวันที่ในไฟล์นี้นับตามเวลาไทยแล้ว"),
    ]:
        line(a, b)

    for ws_ in wb.worksheets:
        ws_.sheet_view.showGridLines = False
        ws_.page_setup.orientation = "landscape"       # พิมพ์ออกมาพอดีความกว้างกระดาษ
        ws_.page_setup.paperSize = ws_.PAPERSIZE_A4
        ws_.page_setup.fitToWidth = 1
        ws_.page_setup.fitToHeight = 0
        ws_.sheet_properties.pageSetUpPr.fitToPage = True
        ws_.print_title_rows = "4:4" if ws_.title != "อ่านก่อน" else None
    wb.save(out_path)
    print(f"เขียนแล้ว: {out_path}  (ตาราง {len(base)} · คอลัมน์ {n_cols} · view {len(views)})")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(ROOT, "deploy", "db_dictionary.xlsx"))
    build(ap.parse_args().out)
